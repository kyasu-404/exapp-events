import re
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from pathlib import PurePosixPath
from uuid import NAMESPACE_URL, uuid5
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook

from ..models import Event, Settings

MONTHS = {
    name.casefold(): i
    for i, name in enumerate(
        [
            "Январь",
            "Февраль",
            "Март",
            "Апрель",
            "Май",
            "Июнь",
            "Июль",
            "Август",
            "Сентябрь",
            "Октябрь",
            "Ноябрь",
            "Декабрь",
        ],
        1,
    )
}
DATE_PATTERN = re.compile(r"^(\d{1,2})\.(\d{1,2})\.?$")
REQUIRED = ("location", "time", "title", "responsible")


def normalized(value) -> str:
    return " ".join(str(value or "").split()).casefold()


def file_context(path: str) -> tuple[int, int | None]:
    parents = list(reversed(PurePosixPath(path).parts[:-1]))
    year = next((int(p) for p in parents if re.fullmatch(r"(19|20)\d{2}", p)), None)
    if year is None:
        raise ValueError("Год не найден в родительских папках")
    month = next((MONTHS[p.strip().casefold()] for p in parents if p.strip().casefold() in MONTHS), None)
    return year, month


def parse_time(value, day: date, settings: Settings) -> tuple[datetime, datetime]:
    if isinstance(value, datetime):
        start, end = value.time(), None
    elif isinstance(value, time):
        start, end = value, None
    else:
        parts = re.split(r"\s*[-–—]\s*", str(value).strip())
        if len(parts) > 2:
            raise ValueError("Некорректное время")
        parsed = []
        for part in parts:
            match = re.fullmatch(r"(\d{1,2})[:.](\d{2})", part.strip())
            if not match:
                raise ValueError("Ожидается время 15:30 или диапазон 15:30–17:00")
            parsed.append(time(int(match[1]), int(match[2])))
        start, end = parsed[0], parsed[1] if len(parsed) == 2 else None
    zone = ZoneInfo(settings.timezone)
    dtstart = datetime.combine(day, start.replace(tzinfo=None), zone)
    dtend = (
        datetime.combine(day, end.replace(tzinfo=None), zone)
        if end
        else dtstart + timedelta(minutes=settings.duration_minutes)
    )
    if end and dtend <= dtstart:
        raise ValueError("Окончание должно быть позже начала в тот же день")
    for value in (dtstart, dtend):
        if value.astimezone(UTC).astimezone(zone).replace(tzinfo=None) != value.replace(tzinfo=None):
            raise ValueError("Время не существует в выбранном часовом поясе")
        if value.replace(fold=0).utcoffset() != value.replace(fold=1).utcoffset():
            raise ValueError("Время неоднозначно при переводе часов")
    return dtstart, dtend


def parse_emails(value) -> tuple[list[str], list[str]]:
    emails, warnings = [], []
    for raw in re.split(r"[;,\n\r]+", str(value or "")):
        raw = raw.strip()
        if not raw:
            continue
        try:
            email = validate_email(raw, check_deliverability=False).normalized
            if email.casefold() not in {e.casefold() for e in emails}:
                emails.append(email)
        except EmailNotValidError:
            warnings.append("Некорректный адрес email пропущен")
    return sorted(emails, key=str.casefold), warnings


def parse_offsets(value) -> tuple[list[int], list[str]]:
    offsets, warnings = set(), []
    for raw in re.split(r"[;,\n]+", str(value or "")):
        raw = raw.strip().casefold()
        if not raw:
            continue
        match = re.fullmatch(r"(\d+)\s*([мmчhдd])", raw)
        if not match or not 0 < int(match[1]) <= 10080:
            warnings.append("Некорректный интервал напоминания пропущен")
            continue
        offsets.add(
            int(match[1]) * {"м": 60, "m": 60, "ч": 3600, "h": 3600, "д": 86400, "d": 86400}[match[2]]
        )
    return sorted(offsets), warnings


@dataclass
class Parsed:
    events: list[Event] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    sheets: list[dict] = field(default_factory=list)
    safe_delete: bool = True
    year: int | None = None

    def issue(self, level, message, sheet="", row=0, protect=False):
        self.issues.append({"level": level, "message": message, "sheet": sheet, "row": row})
        if protect:
            self.safe_delete = False


def parse_xlsx(content: bytes, file_id: str, path: str, settings: Settings) -> Parsed:
    result = Parsed()
    try:
        result.year, folder_month = file_context(path)
        # Bound compressed input and expanded XML. Formatting-only rows are still skipped below.
        if len(content) > 32 * 1024 * 1024:
            raise ValueError("XLSX превышает 32 МБ")
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError("Распакованный XLSX превышает 128 МБ")
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True, keep_links=False)
    except Exception as exc:
        result.issue(
            "Error",
            str(exc) if isinstance(exc, ValueError) else "XLSX не читается; прежние события сохранены",
            protect=True,
        )
        return result
    formula_workbook = None
    formula_rows = {}
    try:
        for sheet in workbook:
            match = DATE_PATTERN.fullmatch(sheet.title.strip())
            if not match:
                result.sheets.append({"name": sheet.title, "status": "ignored"})
                continue
            try:
                day = date(result.year, int(match[2]), int(match[1]))
            except ValueError:
                result.issue("Error", "Недопустимая дата листа", sheet.title, protect=True)
                continue
            if settings.check_month and folder_month and day.month != folder_month:
                result.issue(
                    "Warning", "Месяц листа не совпадает с папкой; лист пропущен", sheet.title, protect=True
                )
                continue
            if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 512:
                result.issue(
                    "Error", "Лист превышает лимит 100000 строк / 512 колонок", sheet.title, protect=True
                )
                continue
            mapping = None
            for row_no, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if mapping is None and row_no > 50:
                    break
                if not any(v is not None and str(v).strip() for v in cells):
                    continue
                if mapping is None:
                    texts = [normalized(c) for c in cells]
                    found = {
                        key: texts.index(normalized(header))
                        for key, header in settings.headers.items()
                        if normalized(header) in texts
                    }
                    if all(key in found for key in REQUIRED):
                        if any(texts.count(normalized(settings.headers[k])) != 1 for k in found):
                            result.issue(
                                "Error", "Повторяющиеся заголовки", sheet.title, row_no, protect=True
                            )
                            break
                        mapping = found
                        for optional in ("emails", "offsets", "public"):
                            if optional not in found:
                                result.issue(
                                    "Warning",
                                    f"Не найден необязательный заголовок: {settings.headers[optional]}",
                                    sheet.title,
                                    row_no,
                                )
                        result.sheets.append(
                            {
                                "name": sheet.title,
                                "date": day.isoformat(),
                                "headers": found,
                                "header_row": row_no,
                            }
                        )
                    elif row_no >= 50:
                        break
                    continue
                values = {key: cells[index] if index < len(cells) else None for key, index in mapping.items()}
                if not any(values.get(k) is not None for k in REQUIRED):
                    continue
                if not any(
                    value is not None and str(value).strip()
                    for key, value in values.items()
                    if key != "location"
                ):
                    # An unused room row is valid. Uncached formulas still protect existing events.
                    if sheet.title not in formula_rows:
                        if formula_workbook is None:
                            formula_workbook = load_workbook(
                                BytesIO(content), data_only=False, read_only=True, keep_links=False
                            )
                        columns = [mapping[key] for key in REQUIRED]
                        first_column, last_column = min(columns), max(columns)
                        formula_rows[sheet.title] = {
                            number
                            for number, raw in enumerate(
                                formula_workbook[sheet.title].iter_rows(
                                    min_row=row_no,
                                    min_col=first_column + 1,
                                    max_col=last_column + 1,
                                ),
                                row_no,
                            )
                            if any(raw[column - first_column].data_type == "f" for column in columns)
                        }
                    if row_no not in formula_rows[sheet.title]:
                        continue
                if not all(values.get(k) is not None and str(values[k]).strip() for k in REQUIRED):
                    missing = [
                        settings.headers[k]
                        for k in REQUIRED
                        if values.get(k) is None or not str(values[k]).strip()
                    ]
                    result.issue(
                        "Warning",
                        "Не заполнены поля: "
                        + ", ".join(missing)
                        + ". Строка пропущена; прежние события файла сохранены.",
                        sheet.title,
                        row_no,
                        protect=True,
                    )
                    continue
                try:
                    start, end = parse_time(values["time"], day, settings)
                except ValueError as exc:
                    result.issue("Error", str(exc), sheet.title, row_no, protect=True)
                    continue
                emails, email_issues = parse_emails(values.get("emails"))
                offsets, offset_issues = parse_offsets(values.get("offsets"))
                for message in email_issues + offset_issues:
                    result.issue("Warning", message, sheet.title, row_no)
                flag = normalized(values.get("public"))
                if flag not in {"", "да", "yes", "true", "1", "+", "нет", "no", "false", "0", "-"}:
                    result.issue(
                        "Warning", "Неизвестное значение «На сайт»: публикация выключена", sheet.title, row_no
                    )
                key = str(
                    uuid5(
                        NAMESPACE_URL,
                        f"exapp_events:{file_id}:{day.isoformat()}:{start.time().isoformat()}:{normalized(values['location'])}",
                    )
                )
                result.events.append(
                    Event(
                        key,
                        str(file_id),
                        sheet.title,
                        start.isoformat(),
                        end.isoformat(),
                        str(values["location"]).strip(),
                        str(values["title"]).strip(),
                        str(values["responsible"]).strip(),
                        emails,
                        offsets,
                        flag in {"да", "yes", "true", "1", "+"},
                    )
                )
            if mapping is None:
                result.issue(
                    "Error",
                    "Обязательные заголовки не найдены в первых 50 строках",
                    sheet.title,
                    protect=True,
                )
        seen, conflicts = set(), set()
        for event in result.events:
            if event.source_key in seen:
                conflicts.add(event.source_key)
            seen.add(event.source_key)
        if conflicts:
            result.issue(
                "Error", f"Конфликт: одинаковые дата, время и место ({len(conflicts)})", protect=True
            )
            result.events = [e for e in result.events if e.source_key not in conflicts]
        if not any("date" in s for s in result.sheets) and not result.issues:
            result.issue("Warning", "Нет листов с датами; удаление прежних событий запрещено", protect=True)
    finally:
        workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
    return result
